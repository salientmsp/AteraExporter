"""
Export Utilities Module
Handles exporting data to various formats (CSV, JSON, Excel)
"""

import csv
import json
import os
from datetime import datetime
import logging

logger = logging.getLogger(__name__)

# Create exports directory if it doesn't exist
EXPORT_DIR = 'exports'
if not os.path.exists(EXPORT_DIR):
    os.makedirs(EXPORT_DIR)


def generate_filename(data_type, export_format):
    """
    Generate a filename for the export

    Args:
        data_type: Type of data being exported (e.g., 'customers', 'tickets')
        export_format: Format of export ('csv', 'json', 'excel')

    Returns:
        Full file path
    """
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{data_type}_{timestamp}.{export_format}"

    if export_format == 'excel':
        filename = f"{data_type}_{timestamp}.xlsx"

    return os.path.join(EXPORT_DIR, filename)


def export_to_csv(data, filename, fieldnames=None):
    """
    Export data to CSV format

    Args:
        data: List of dictionaries to export
        filename: Output filename
        fieldnames: Optional list of field names (columns)

    Returns:
        True on success, False on failure
    """
    try:
        if not data:
            logger.warning("No data to export to CSV")
            return False

        # If fieldnames not provided, use keys from first item
        if fieldnames is None:
            fieldnames = list(data[0].keys())

        with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(data)

        logger.info(f"Successfully exported {len(data)} records to CSV: {filename}")
        return True

    except Exception as e:
        logger.error(f"Error exporting to CSV: {str(e)}")
        return False


def export_to_json(data, filename, indent=2):
    """
    Export data to JSON format

    Args:
        data: Data to export (list or dict)
        filename: Output filename
        indent: JSON indentation level

    Returns:
        True on success, False on failure
    """
    try:
        if not data:
            logger.warning("No data to export to JSON")
            return False

        with open(filename, 'w', encoding='utf-8') as jsonfile:
            json.dump(data, jsonfile, indent=indent, default=str)

        logger.info(f"Successfully exported data to JSON: {filename}")
        return True

    except Exception as e:
        logger.error(f"Error exporting to JSON: {str(e)}")
        return False


def export_to_excel(data, filename, sheet_name='Data'):
    """
    Export data to Excel format

    Args:
        data: List of dictionaries to export
        filename: Output filename
        sheet_name: Name of the Excel sheet

    Returns:
        True on success, False on failure
    """
    try:
        # Import openpyxl for Excel support
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            return False

        if not data:
            logger.warning("No data to export to Excel")
            return False

        # Create workbook and sheet
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Get fieldnames from first item
        fieldnames = list(data[0].keys())

        # Write header row with formatting
        for col_num, fieldname in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = fieldname
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')

        # Write data rows
        for row_num, item in enumerate(data, 2):
            for col_num, fieldname in enumerate(fieldnames, 1):
                value = item.get(fieldname)
                # Convert datetime objects to strings
                if isinstance(value, datetime):
                    value = value.strftime('%Y-%m-%d %H:%M:%S')
                ws.cell(row=row_num, column=col_num, value=value)

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save workbook
        wb.save(filename)

        logger.info(f"Successfully exported {len(data)} records to Excel: {filename}")
        return True

    except Exception as e:
        logger.error(f"Error exporting to Excel: {str(e)}")
        return False


def model_to_dict(model_instance, exclude_fields=None):
    """
    Convert SQLAlchemy model instance to dictionary

    Args:
        model_instance: SQLAlchemy model instance
        exclude_fields: List of field names to exclude

    Returns:
        Dictionary representation of model
    """
    if exclude_fields is None:
        exclude_fields = []

    data = {}
    for column in model_instance.__table__.columns:
        if column.name not in exclude_fields:
            value = getattr(model_instance, column.name)
            # Convert datetime to string
            if isinstance(value, datetime):
                value = value.strftime('%Y-%m-%d %H:%M:%S')
            data[column.name] = value

    return data


def models_to_list(model_instances, exclude_fields=None):
    """
    Convert list of SQLAlchemy model instances to list of dictionaries

    Args:
        model_instances: List of SQLAlchemy model instances
        exclude_fields: List of field names to exclude

    Returns:
        List of dictionaries
    """
    return [model_to_dict(instance, exclude_fields) for instance in model_instances]


def export_models(model_instances, data_type, export_format, exclude_fields=None):
    """
    Export SQLAlchemy models to specified format

    Args:
        model_instances: List of SQLAlchemy model instances
        data_type: Type of data (for filename)
        export_format: Format to export to ('csv', 'json', 'excel')
        exclude_fields: Fields to exclude from export

    Returns:
        Tuple of (success, filename, record_count, error_message)
    """
    try:
        if not model_instances:
            return False, None, 0, "No data to export"

        # Convert models to dictionaries
        data = models_to_list(model_instances, exclude_fields)

        # Generate filename
        filename = generate_filename(data_type, export_format)

        # Export based on format
        if export_format == 'csv':
            success = export_to_csv(data, filename)
        elif export_format == 'json':
            success = export_to_json(data, filename)
        elif export_format == 'excel':
            success = export_to_excel(data, filename)
        else:
            return False, None, 0, f"Unsupported export format: {export_format}"

        if success:
            return True, filename, len(data), None
        else:
            return False, None, 0, "Export failed"

    except Exception as e:
        logger.error(f"Error in export_models: {str(e)}")
        return False, None, 0, str(e)
